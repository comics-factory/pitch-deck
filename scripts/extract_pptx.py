#!/usr/bin/env python3
"""Extract text and images from PowerPoint file."""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import json

def extract_pptx(pptx_path: str, output_dir: str):
    """Extract all text and images from PPTX."""
    pptx_file = Path(pptx_path)
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    
    prs = Presentation(pptx_file)
    
    slides_data = []
    image_counter = 0
    
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_info = {
            "slide_number": slide_num,
            "title": "",
            "content": [],
            "images": [],
            "notes": ""
        }
        
        # Extract text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    # Try to identify title
                    if hasattr(shape, "is_placeholder") and shape.is_placeholder:
                        if shape.placeholder_format.type == 1:  # Title placeholder
                            slide_info["title"] = text
                        else:
                            slide_info["content"].append(text)
                    else:
                        if not slide_info["title"] and len(text) < 100:
                            slide_info["title"] = text
                        else:
                            slide_info["content"].append(text)
            
            # Extract images
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                image_counter += 1
                image_filename = f"slide{slide_num:02d}_img{image_counter:02d}.{image.ext}"
                image_path = output_path / image_filename
                
                with open(image_path, "wb") as f:
                    f.write(image.blob)
                
                slide_info["images"].append(image_filename)
                print(f"  Extracted: {image_filename}")
        
        # Extract notes
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame:
                slide_info["notes"] = notes_frame.text.strip()
        
        slides_data.append(slide_info)
        print(f"Slide {slide_num}: {slide_info['title'][:50] if slide_info['title'] else '(no title)'}")
    
    # Save as JSON
    json_path = output_path / "extracted_content.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(slides_data, f, indent=2, ensure_ascii=False)
    
    print(f"\n✓ Extracted {len(slides_data)} slides to {output_path}")
    print(f"✓ Content saved to {json_path}")
    return slides_data

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: uv run python scripts/extract_pptx.py INPUT.pptx OUTPUT_DIR")
        sys.exit(1)
    
    extract_pptx(sys.argv[1], sys.argv[2])
